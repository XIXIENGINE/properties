#!/usr/bin/env python3
"""
GitHub Actions 러너용 자산 브리핑 생성/발송기.

동작:
  1. 냥돌폴리오 앱의 보유현황 CSV를 받는다 (기본 출처).
     앱이 이미 거래를 접어 보유수량·평단을 내고 실시간 시세·환율까지 반영하므로
     여기서 다시 계산하지 않는다. 계산 경로가 둘이면 화면과 브리핑이 갈린다.
  2. 전일 스냅샷과 비교해 브리핑을 만든다(brief.py의 계산/렌더러 재사용).
  3. Slack Incoming Webhook으로 Block Kit 메시지를 보낸다.
  4. 스냅샷/히스토리/HTML을 저장한다.

앱에서 받지 못하면(Supabase 일시정지 등) 조용히 죽지 않고 원인을 Slack으로 알린다.
구글 시트 파싱 경로는 --source sheet 로 남겨 두었다 (앱이 오래 내려갔을 때의 비상용).

환경변수:
  SLACK_WEBHOOK_URL      : Slack Incoming Webhook (없으면 발송 생략, 미리보기만 출력)
  NYANGDOL_APP_URL       : 앱 주소 (예: https://nyangdol-folio-969u.vercel.app)
  NYANGDOL_APP_PASSWORD  : 앱 공용 비밀번호 (앱의 APP_PASSWORD 와 같은 값)
  BRIEF_SOURCE           : app(기본) | sheet
"""

import argparse
import json
import os
import re
import sys
import tempfile
import urllib.request
from datetime import datetime

import app_source  # 냥돌폴리오 앱 CSV 내보내기
import brief  # 같은 저장소의 파싱/집계/렌더러
import prices  # 실시간 시세 조회

SHEET_ID = "1MLuYEdJUUuhQ6V3hp7WSmsIFeFRo9LIHL23_wcdezFw"
XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

# 시트 헤더명 → 내부 필드
HEADER_MAP = {
    "증권사": "broker", "계좌": "account", "종목명": "name", "티커": "ticker",
    "수량": "qty", "매입금액": "cost", "평가금액": "eval", "성격": "nature",
    "No.": "no", "no": "no", "No": "no",
}


def download_xlsx(dest):
    req = urllib.request.Request(XLSX_URL, headers={"User-Agent": "asset-brief/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    with open(dest, "wb") as f:
        f.write(data)
    return len(data)


def _cell_str(v):
    return "" if v is None else str(v).strip()


def _to_int(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().replace(",", "").replace("원", "")
    if s in ("", "-"):
        return 0
    try:
        return int(round(float(s)))
    except ValueError:
        return 0


def _sheet_owner(rows, header_idx, fallback):
    """입력 탭의 제목 행에서 보유자 이름을 뽑는다.

    예: '입력 시트 · 이수진   (증권사 → 계좌 순 · 이수진 자산만)' -> '이수진'
    """
    for row in rows[:header_idx]:
        for c in row:
            t = _cell_str(c)
            if "입력" in t and "시트" in t:
                m = re.search(r"입력\s*시트\s*[·:\-]\s*([^\s(·]+)", t)
                if m:
                    return m.group(1).strip()
    # 폴백: 시트 이름에서 구분자 뒤쪽('입력_이수진' 등)
    t = (fallback or "").strip()
    for sep in ("·", "_", "-", " "):
        if sep in t:
            tail = t.split(sep)[-1].strip()
            if tail and "입력" not in tail:
                return tail
    return t or "미상"


def parse_xlsx(path):
    """모든 '입력 탭'을 읽어 보유자별 보유 종목을 합쳐서 돌려준다.

    입력 탭 판별: 헤더에 '증권사'·'종목명'·'평가금액'이 함께 있는 시트.
    (보유자 열을 쓰는 '상세현황' 탭이나 대시보드·분류 탭은 자동으로 제외된다.)
    같은 계좌·종목명이 두 사람 모두에게 있을 수 있으므로 key에 보유자를 포함한다.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    merged, sheets = [], []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_idx, colmap = None, {}
        for i, row in enumerate(rows):
            labels = [_cell_str(c) for c in row]
            if "종목명" in labels and "평가금액" in labels and "증권사" in labels:
                header_idx = i
                for j, lab in enumerate(labels):
                    if lab in HEADER_MAP:
                        colmap[HEADER_MAP[lab]] = j
                break
        if header_idx is None or "name" not in colmap or "eval" not in colmap:
            continue

        owner = _sheet_owner(rows, header_idx, getattr(ws, "title", ""))
        holdings = []
        for row in rows[header_idx + 1:]:
            def cell(field):
                j = colmap.get(field)
                return row[j] if (j is not None and j < len(row)) else None
            no_raw = _cell_str(cell("no"))
            name = _cell_str(cell("name"))
            # No.가 정수인 행만 유효한 보유 종목
            if not no_raw.replace(".0", "").isdigit():
                continue
            if not name:
                continue
            broker = _cell_str(cell("broker"))
            account = _cell_str(cell("account"))
            holdings.append({
                "no": int(float(no_raw)),
                "owner": owner,
                "broker": broker,
                "account": account,
                "name": name,
                "ticker": _cell_str(cell("ticker")),
                "qty": _cell_str(cell("qty")),
                "cost": _to_int(cell("cost")),
                "eval": _to_int(cell("eval")),
                "ret": None,
                "nature": _cell_str(cell("nature")) or "기타",
                "key": f"{owner}|{broker}|{account}|{name}",
            })
        if holdings:
            merged.extend(holdings)
            sheets.append(f"{owner}({len(holdings)})")
    wb.close()
    if sheets:
        print(f"[parse] 입력 탭 {len(sheets)}개: {', '.join(sheets)}")
    return _merge_duplicates(merged)


def _merge_duplicates(holdings):
    """같은 보유자·계좌·종목이 여러 줄(분할 매수 등)이면 한 줄로 합친다.

    key가 겹치면 스냅샷 딕셔너리에서 서로 덮어써져 종목별 전일 대비가 틀어지므로,
    수량·매입금액·평가금액을 합산해 하나의 보유로 만든다.
    """
    out, index = [], {}
    for h in holdings:
        k = h["key"]
        if k not in index:
            index[k] = len(out)
            out.append(h)
            continue
        tgt = out[index[k]]
        tgt["cost"] += h["cost"]
        tgt["eval"] += h["eval"]
        try:
            tgt["qty"] = str(float(tgt["qty"] or 0) + float(h["qty"] or 0))
        except ValueError:
            pass
    if len(out) != len(holdings):
        print(f"[parse] 중복 종목 {len(holdings) - len(out)}건 합산 "
              f"(같은 계좌·종목의 여러 줄을 1건으로)")
    return out


def post_slack(webhook, blocks, fallback_text):
    body = {"text": fallback_text}
    # blocks 를 null 로 보내면 Slack 이 거부한다. 실패 알림처럼 텍스트만 보낼 때가 있다.
    if blocks:
        body["blocks"] = blocks
    payload = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        webhook, data=payload,
        headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.status, r.read().decode("utf-8", "replace")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--state", default="data/latest.json")
    ap.add_argument("--history", default="data/history.csv")
    ap.add_argument("--out-md", default="data/brief.md")
    ap.add_argument("--out-html", default="data/brief.html")
    ap.add_argument("--snapshots", default="data/snapshots")
    ap.add_argument("--date", default=None, help="기준일 YYYY-MM-DD (기본: 오늘 KST)")
    ap.add_argument("--dry-run", action="store_true", help="스냅샷/히스토리 미저장")
    ap.add_argument("--no-slack", action="store_true", help="Slack 발송 생략")
    ap.add_argument("--always-send", action="store_true",
                    help="직전과 동일해도 강제 발송(수동 테스트용). 환경변수 ALWAYS_SEND로도 지정")
    ap.add_argument("--live-prices", action="store_true",
                    help="실시간 시세로 평가금액 자동 계산(수량×현재가×환율). 환경변수 LIVE_PRICES로도")
    ap.add_argument("--input", default=None,
                    help="xlsx 대신 로컬 마크다운 파일로 파싱(테스트용)")
    ap.add_argument("--source", choices=("app", "sheet"), default=None,
                    help="데이터 출처. app=냥돌폴리오 앱(기본) · sheet=구글 시트(비상용). "
                         "환경변수 BRIEF_SOURCE 로도 지정")
    args = ap.parse_args()

    if args.date:
        date_kst = datetime.strptime(args.date, "%Y-%m-%d").replace(tzinfo=brief.KST)
    else:
        date_kst = datetime.now(brief.KST)
    date_str = date_kst.strftime("%Y-%m-%d")

    source = args.source or os.environ.get("BRIEF_SOURCE", "").strip() or "app"
    webhook = os.environ.get("SLACK_WEBHOOK_URL", "").strip()

    # 1) 보유 종목 파싱
    from_app = False
    if args.input:
        holdings = brief.parse_holdings(open(args.input, encoding="utf-8").read())
        print(f"[in] 로컬 마크다운 파싱: {args.input}")
    elif source == "app":
        # 앱이 이미 시세·환율까지 반영해 계산한 값을 받는다. 여기서 다시 계산하지 않는다.
        try:
            base_url, password = app_source.config_from_env()
            holdings = app_source.load_holdings(base_url, password)
        except app_source.AppUnavailable as e:
            # 조용히 죽으면 며칠 뒤에야 안 온 걸 알게 된다. 원인을 Slack 으로 알린다.
            print(f"ERROR: {e.message} {e.hint or ''}".strip(), file=sys.stderr)
            if webhook and not args.no_slack:
                brief_text = f"⚠️ 자산 브리핑을 만들지 못했습니다\n{e.message}"
                if e.hint:
                    brief_text += f"\n{e.hint}"
                post_slack(webhook, None, brief_text)
                print("[slack] 실패 알림 발송")
            sys.exit(2)
        from_app = True
        print(f"[fetch] 앱 {base_url} · 보유현황 CSV")
    else:
        tmp = os.path.join(tempfile.gettempdir(), "sheet.xlsx")
        n = download_xlsx(tmp)
        print(f"[fetch] xlsx {n:,} bytes")
        holdings = parse_xlsx(tmp)
    if not holdings:
        print("ERROR: 보유 종목을 파싱하지 못했습니다.", file=sys.stderr)
        sys.exit(2)
    print(f"[parse] {len(holdings)}개 종목 (출처: {'앱' if from_app else '시트'})")

    # 실시간 시세 반영(옵션)
    # 앱에서 받은 값은 이미 앱이 시세·환율을 반영한 결과다. 여기서 또 조회하면
    # 같은 자산이 두 경로로 계산되어 화면과 브리핑이 갈린다.
    env_live = os.environ.get("LIVE_PRICES", "").strip().lower() not in ("", "0", "false", "no")
    extra_note = None
    if from_app:
        if args.live_prices or env_live:
            print("[price] 앱 소스에서는 시세를 다시 조회하지 않습니다 (앱이 이미 반영).")
        extra_note = "냥돌폴리오 앱 기준 (시세·환율 앱 계산값)"
    elif args.live_prices or env_live:
        holdings, rep = prices.enrich_holdings(holdings)
        extra_note = prices.report_note(rep)
        fx_txt = f"{rep['fx']:,.1f}" if rep.get("fx") else "N/A"
        print(f"[price] 실시간 {rep['live']}건 · 폴백 {rep['fallback']}건 · "
              f"보유0 스킵 {rep['held_skipped']}건 · USD/KRW {fx_txt}")
        print(f"[price] 시트합계 {rep.get('sheet_total', 0):,}원 → "
              f"시세합계 {rep.get('live_total', 0):,}원 "
              f"(차이 {rep.get('live_total', 0) - rep.get('sheet_total', 0):+,}원)")
        if rep["failures"]:
            print(f"[price] 조회 실패(시트값 사용): {', '.join(rep['failures'])}")
        if os.environ.get("PRICE_DEBUG", "").strip() and rep.get("debug"):
            for name, qty, price, se, ne in rep["debug"][:8]:
                print(f"[debug] {name}: {qty:g}주 × {price:,.2f} = {ne:,} (시트 {se:,})")

    agg = brief.aggregate(holdings)
    print(f"[agg] 총 평가금액 {agg['total_eval']:,}원 · 매입 {agg['total_cost']:,}원")

    prev = None
    if os.path.exists(args.state):
        try:
            prev = json.load(open(args.state, encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            prev = None
    print(f"[prev] {'직전 스냅샷 '+prev['date'] if prev else '없음(첫 실행)'}")

    # 키 체계가 바뀐 경우(예: 보유자 구분 추가) 이전 스냅샷과는 비교할 수 없다.
    # 잘못된 급등락을 보여주는 대신 오늘을 새 기준으로 삼는다.
    baseline_reset = False
    if prev and prev.get("holdings"):
        if not (set(prev["holdings"]) & {h["key"] for h in holdings}):
            print("[prev] ⚠️ 이전 스냅샷과 종목 키 체계가 달라 비교 기준을 재설정합니다 "
                  "(오늘이 새 기준 · 전일 대비는 다음 실행부터)")
            prev = None
            baseline_reset = True

    # 주말·휴일 등 직전과 완전히 동일하면 무의미한 알림을 피한다.
    env_always = os.environ.get("ALWAYS_SEND", "").strip().lower() not in ("", "0", "false", "no")
    always_send = args.always_send or env_always
    unchanged = (prev is not None
                 and prev.get("total_eval") == agg["total_eval"]
                 and prev.get("holdings") == agg["holdings"])
    if unchanged and not always_send:
        print("[skip] 직전 스냅샷과 자산 구성·평가금액이 동일 → 발송/저장 생략 "
              "(주말·휴일 등 변동 없음). 강제 발송은 --always-send / ALWAYS_SEND=1")
        return

    if baseline_reset:
        note = "가계 합산(보유자 구분) 기준으로 변경 — 전일 대비는 다음 브리핑부터"
        extra_note = f"{extra_note} · {note}" if extra_note else note

    summary = brief.compute_summary(holdings, agg, prev, date_kst, extra_note=extra_note)
    md, acct = brief.build_message(holdings, agg, prev, date_kst, extra_note=extra_note)
    blocks = brief.build_slack_blocks(summary)
    html = brief.build_html(summary)

    os.makedirs(os.path.dirname(args.out_md) or ".", exist_ok=True)
    open(args.out_md, "w", encoding="utf-8").write(md)
    open(args.out_html, "w", encoding="utf-8").write(html)
    print(f"[render] {args.out_md} · {args.out_html}")

    # 4) Slack 발송
    if args.no_slack or not webhook:
        reason = "--no-slack" if args.no_slack else "SLACK_WEBHOOK_URL 없음"
        print(f"[slack] 발송 생략({reason}). 블록 미리보기:")
        print(json.dumps({"blocks": blocks}, ensure_ascii=False)[:1200])
    else:
        fallback = f"자산 데일리 브리핑 {date_str} · 총 {agg['total_eval']:,}원"
        status, body = post_slack(webhook, blocks, fallback)
        print(f"[slack] HTTP {status} {body[:200]}")
        if status != 200:
            sys.exit(3)

    # 5) 스냅샷 저장
    if args.dry_run:
        print("[persist] --dry-run: 저장 생략")
    else:
        st = brief.persist_snapshot(agg, acct, prev, date_str,
                                    args.state, args.history, args.snapshots)
        print(f"[persist] 스냅샷 저장: {st['date']}")


if __name__ == "__main__":
    main()
